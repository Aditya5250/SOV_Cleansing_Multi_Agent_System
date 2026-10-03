import os
import openpyxl
import pytest
from backend.app.orchestration.orchestrator import orchestrator
from backend.app.orchestration.state import state_manager
from backend.app.schemas.sov_schema import TARGET_FIELDS
from backend.app.services.validation_service import ValidationService

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "sample_data")

@pytest.mark.parametrize("filename", [
    "Sample_SOV_1_Standard.xlsx",
    "Sample_SOV_2_Complex_Semantic.xlsx",
    "Sample_SOV_3_MultiSheet_Messy.xlsx"
])
def test_end_to_end_pipeline_all_samples(filename):
    file_path = os.path.join(SAMPLE_DIR, filename)
    assert os.path.exists(file_path), f"File {filename} does not exist"

    # Step 1: Initialize session
    file_info = {"filename": filename, "extension": ".xlsx", "size_kb": 15.0}
    state = state_manager.create_session(file_info=file_info, raw_file_path=file_path)

    # Step 2: Run Intake Pipeline (Agents 1, 2, 3)
    updated_state = orchestrator.run_intake_pipeline(state.session_id)

    # Assertions on Intake
    assert updated_state.selected_sheet is not None
    assert updated_state.schema_mapping is not None
    assert updated_state.data_quality is not None

    # Step 3: Human Review Simulation: Approve all high-confidence recommendations
    decisions = {}
    for rec in updated_state.recommendations:
        if rec.confidence >= 0.85:
            decisions[rec.id] = {"status": "accepted"}
        else:
            decisions[rec.id] = {"status": "rejected"}

    # Step 4: Execute Controlled Transformation (Agent 4)
    result = orchestrator.execute_controlled_transformation(
        session_id=state.session_id,
        decisions=decisions,
        approved_by="Automated Test Runner"
    )

    assert result.success is True
    assert result.column_count == 17
    assert result.columns == TARGET_FIELDS
    assert result.audit_log_count >= 1

    # Step 5: Verify exported Cleaned_SOV.xlsx
    sov_path = updated_state.cleaned_file_path
    assert os.path.exists(sov_path)

    wb = openpyxl.load_workbook(sov_path)
    assert "Cleaned_SOV" in wb.sheetnames
    ws = wb["Cleaned_SOV"]

    # Verify Header row
    actual_headers = [cell.value for cell in ws[1]]
    assert actual_headers == TARGET_FIELDS
    assert len(ws.merged_cells.ranges) == 0  # No merged cells
    wb.close()

    # Step 6: Verify exported Audit_Log.xlsx
    audit_path = updated_state.audit_xlsx_path
    assert os.path.exists(audit_path)
    wb_audit = openpyxl.load_workbook(audit_path)
    assert "Audit_Log" in wb_audit.sheetnames
    wb_audit.close()
